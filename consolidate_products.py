"""Merge every scraped output file into one unified product list."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment

COLUMNS = [
    "product_name", "brand", "description", "manufacturer_name", "importer_name",
    "country_of_manufacture", "company_name", "company_profile", "phone", "email",
    "company_address", "website", "brochure_links", "machine_usage", "tests_performed",
    "used_in_which_test", "product_url", "source_website", "data_source",
]

# (file path, data source label) for the venushealthcare/tradeindia-style schema
NAME_SCHEMA_SOURCES = [
    ("products.json", "Venus Health Care"),
    ("tradeindia_products.json", "TradeIndia clinical analyzers"),
    ("tradeindia_kanad.json", "TradeIndia clinical analyzers"),
    ("standard-f200.json", "Venus Health Care"),
]

# (directory, data source label) for the medical_equipment_scraper.py schema
SOURCE_OUTPUT_DIRS = [
    ("source_outputs", "TradeIndia listings"),
    ("source_outputs_manufacturers", "Indian IVD manufacturers"),
]


def joined(value: Any) -> str:
    if isinstance(value, list):
        return " | ".join(str(item) for item in value if item)
    return str(value or "")


def from_name_schema(product: dict[str, Any], label: str) -> dict[str, str]:
    return {
        "product_name": product.get("name", ""),
        "brand": product.get("brand", ""),
        "description": product.get("description", ""),
        "manufacturer_name": product.get("manufacturer_company", ""),
        "importer_name": product.get("importer_company", ""),
        "country_of_manufacture": product.get("country_of_manufacture", ""),
        "company_name": product.get("company_name") or product.get("supplier_company", ""),
        "company_profile": product.get("company_profile", ""),
        "phone": product.get("phone") or product.get("seller_phone", ""),
        "email": product.get("email", ""),
        "company_address": product.get("company_address", ""),
        "website": product.get("website", ""),
        "brochure_links": joined(product.get("brochure_links")),
        "machine_usage": product.get("machine_usage", ""),
        "tests_performed": joined(product.get("tests_performed")),
        "used_in_which_test": "",
        "product_url": product.get("url", ""),
        "source_website": "",
        "data_source": label,
    }


def from_source_output_schema(product: dict[str, Any], label: str) -> dict[str, str]:
    return {
        "product_name": product.get("machine_product_name", ""),
        "brand": "",
        "description": "",
        "manufacturer_name": product.get("manufacturer_name", ""),
        "importer_name": product.get("importer_name", ""),
        "country_of_manufacture": "",
        "company_name": product.get("company_name", ""),
        "company_profile": product.get("company_profile_description", ""),
        "phone": product.get("company_phone_number", ""),
        "email": product.get("company_email_id", ""),
        "company_address": product.get("company_location_address", ""),
        "website": "",
        "brochure_links": product.get("product_brochure_link", ""),
        "machine_usage": "",
        "tests_performed": "",
        "used_in_which_test": product.get("used_in_which_test", ""),
        "product_url": product.get("product_source_link", ""),
        "source_website": product.get("source_website", ""),
        "data_source": label,
    }


def load_products() -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for filename, label in NAME_SCHEMA_SOURCES:
        path = Path(filename)
        if not path.exists():
            continue
        data = json.loads(path.read_text(encoding="utf-8"))
        products = data.get("products", data if isinstance(data, list) else [])
        rows.extend(from_name_schema(product, label) for product in products)
    for dirname, label in SOURCE_OUTPUT_DIRS:
        directory = Path(dirname)
        if not directory.is_dir():
            continue
        for path in sorted(directory.glob("source_*.json")):
            data = json.loads(path.read_text(encoding="utf-8"))
            rows.extend(from_source_output_schema(product, label) for product in data.get("products", []))
    return rows


def deduplicate(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    seen: set[tuple[str, str]] = set()
    result = []
    for row in rows:
        key = (row["product_url"], row["product_name"].strip().lower())
        if key in seen:
            continue
        seen.add(key)
        result.append(row)
    return result


def save_excel(rows: list[dict[str, str]], output_path: str) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append(COLUMNS)
    for row in rows:
        sheet.append([row.get(column, "") for column in COLUMNS])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.iter_cols(1, len(COLUMNS)):
        sheet.column_dimensions[column[0].column_letter].width = 32
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(output_path)


def main() -> None:
    rows = deduplicate(load_products())
    Path("all_products.json").write_text(
        json.dumps({"products": rows, "count": len(rows)}, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    save_excel(rows, "all_products.xlsx")
    by_source: dict[str, int] = {}
    for row in rows:
        by_source[row["data_source"]] = by_source.get(row["data_source"], 0) + 1
    print(f"Consolidated {len(rows)} unique products into all_products.json / all_products.xlsx")
    for source, count in sorted(by_source.items(), key=lambda item: -item[1]):
        print(f"  {source}: {count}")


if __name__ == "__main__":
    main()
