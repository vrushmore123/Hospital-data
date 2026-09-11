"""Scrape TradeIndia clinical-analyzer products into the shared product schema."""

from __future__ import annotations

import argparse
import json
import re
import time
from typing import Any
from urllib.parse import parse_qsl, urlencode, urljoin, urlsplit, urlunsplit

import requests
from bs4 import BeautifulSoup, Tag
from openpyxl import Workbook
from openpyxl.styles import Alignment

SEARCH_URL = "https://www.tradeindia.com/search.html?keyword=clinical+analyzer"
DEFAULT_URL = "https://www.tradeindia.com/products/6p-kanad-photon-clinical-chemistry-analyzer-c6374629.html"
PRODUCT_ROOT = "https://www.tradeindia.com/products/"
HEADERS = {"User-Agent": "Mozilla/5.0 (compatible; ProductDataExtractor/1.0)"}
COLUMNS = [
    "name", "brand", "description", "url", "image", "features", "applications", "related_products",
    "manufacturer_company", "importer_company", "country_of_manufacture", "supplier_company",
    "company_name", "company_profile", "phone", "email", "website", "company_address",
    "brochure_links", "machine_usage", "tests_performed", "seller_phone",
]


def clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def unique(values: list[str]) -> list[str]:
    return list(dict.fromkeys(value for value in values if value))


def canonical_url(url: str) -> str:
    parts = urlsplit(url)
    query = [(key, value) for key, value in parse_qsl(parts.query) if key.lower() not in {"utm_source", "utm_medium", "utm_campaign"}]
    return urlunsplit((parts.scheme or "https", parts.netloc, parts.path.rstrip("/"), urlencode(query), ""))


def metadata(soup: BeautifulSoup, selector: str, attribute: str = "content") -> str:
    node = soup.select_one(selector)
    return clean_text(node.get(attribute, "")) if node else ""


def json_ld_objects(soup: BeautifulSoup) -> list[dict[str, Any]]:
    objects = []
    for script in soup.find_all("script", attrs={"type": "application/ld+json"}):
        try:
            data = json.loads(script.string or script.get_text())
        except (json.JSONDecodeError, TypeError):
            continue
        candidates = data if isinstance(data, list) else data.get("@graph", [data]) if isinstance(data, dict) else []
        objects.extend(candidate for candidate in candidates if isinstance(candidate, dict))
    return objects


def first_json_ld(soup: BeautifulSoup, types: tuple[str, ...]) -> dict[str, Any]:
    for item in json_ld_objects(soup):
        item_types = item.get("@type", [])
        item_types = item_types if isinstance(item_types, list) else [item_types]
        if any(item_type in types for item_type in item_types):
            return item
    return {}


def text_after_label(soup: BeautifulSoup, labels: tuple[str, ...]) -> str:
    pattern = re.compile(r"^(?:" + "|".join(re.escape(label) for label in labels) + r")\s*:?$", re.IGNORECASE)
    for node in soup.find_all(string=pattern):
        parent = node.parent
        if not parent:
            continue
        sibling = parent.find_next_sibling()
        if sibling:
            value = clean_text(sibling.get_text(" "))
            if value:
                return value
        text = clean_text(parent.parent.get_text(" ")) if parent.parent else ""
        value = re.sub(pattern, "", text).strip(" :|-\t")
        if value:
            return value
    return ""


def next_data_company(soup: BeautifulSoup) -> dict[str, str]:
    """Read seller details TradeIndia embeds as page state (__NEXT_DATA__) rather than visible text."""
    script = soup.find("script", id="__NEXT_DATA__", type="application/json")
    if not script or not script.string:
        return {}
    try:
        data = json.loads(script.string)
        details = data["props"]["pageProps"]["initialState"]["product"]["PDP_page"]["PDP_page_res"]["company_details"]
    except (json.JSONDecodeError, KeyError, TypeError):
        return {}
    profile_html = details.get("company_details") or ""
    return {
        "company_name": clean_text(str(details.get("co_name") or "")),
        "company_profile": clean_text(BeautifulSoup(profile_html, "html.parser").get_text(" ")),
        "company_address": clean_text(str(details.get("address") or "")),
        "website": urljoin("https://www.tradeindia.com", details.get("profile_url") or ""),
    }


def company_profile(soup: BeautifulSoup, product: dict[str, Any], page_url: str) -> dict[str, str]:
    seller = product.get("seller", {})
    seller_name = seller.get("name", "") if isinstance(seller, dict) else str(seller)
    body_text = clean_text(soup.get_text(" "))
    profile_match = re.search(r"Company Details\s+(.*?)(?:Most Popular Products|Contact Seller|Send Inquiry|Explore Related)", body_text, re.IGNORECASE)
    profile = clean_text(profile_match.group(1)) if profile_match else ""
    name = clean_text(str(seller_name or text_after_label(soup, ("Seller", "Company Name"))))
    phone = text_after_label(soup, ("Phone", "Mobile", "Contact Number"))
    email = text_after_label(soup, ("Email", "Email Id"))
    website = text_after_label(soup, ("Website",))
    address = text_after_label(soup, ("Address", "Location"))
    result = {
        "company_name": name,
        "company_profile": profile,
        "phone": phone,
        "email": email,
        "website": website or (urljoin(page_url, seller.get("url", "")) if isinstance(seller, dict) else ""),
        "company_address": address,
    }
    for field, value in next_data_company(soup).items():
        if value and not result[field]:
            result[field] = value
    return result


def related_products(soup: BeautifulSoup, page_url: str) -> list[dict[str, str]]:
    results = []
    for link in soup.find_all("a", href=True):
        href = canonical_url(urljoin(page_url, link["href"]))
        name = clean_text(link.get_text(" "))
        if href.startswith(PRODUCT_ROOT) and href != canonical_url(page_url) and name:
            results.append({"name": name, "url": href})
    seen = set()
    unique_results = []
    for item in results:
        if item["url"] not in seen:
            seen.add(item["url"])
            unique_results.append(item)
    return unique_results


def product_links(soup: BeautifulSoup, page_url: str) -> list[str]:
    links = []
    for link in soup.find_all("a", href=True):
        href = canonical_url(urljoin(page_url, link["href"]))
        if href.startswith(PRODUCT_ROOT) and re.search(r"-c\d+\.html$", href):
            links.append(href)
    return list(dict.fromkeys(links))


def search_page_links(soup: BeautifulSoup, page_url: str) -> list[str]:
    """Find paginated search-result URLs while ignoring product links."""
    links = []
    for link in soup.find_all("a", href=True):
        href = canonical_url(urljoin(page_url, link["href"]))
        parsed = urlsplit(href)
        if parsed.path.rstrip("/") == "/search.html" and any(key == "page" for key, _ in parse_qsl(parsed.query)):
            links.append(href)
    return list(dict.fromkeys(links))


def extract_product(url: str, timeout: int = 30) -> dict[str, Any]:
    response = requests.get(url, headers=HEADERS, timeout=timeout)
    response.raise_for_status()
    page_url = canonical_url(response.url)
    soup = BeautifulSoup(response.text, "html.parser")
    structured = first_json_ld(soup, ("Product",))
    offer = structured.get("offers", {}) if isinstance(structured.get("offers"), dict) else {}
    seller = structured.get("brand", {}) or structured.get("seller", {})
    seller_name = seller.get("name", "") if isinstance(seller, dict) else str(seller)
    heading = soup.find(["h1", "h2"], string=lambda value: value and "clinical chemistry analyzer" in value.lower())
    title = clean_text(heading.get_text(" ")) if heading else clean_text(str(structured.get("name", "")))
    title = title or metadata(soup, 'meta[property="og:title"]')
    title = re.split(r"\s+at\s+\d[\d,.]*\s+INR\b", title, maxsplit=1, flags=re.IGNORECASE)[0].strip()
    description = clean_text(str(structured.get("description", ""))) or metadata(soup, 'meta[name="description"]')
    image = structured.get("image", "") or metadata(soup, 'meta[property="og:image"]')
    image = urljoin(page_url, image[0] if isinstance(image, list) and image else str(image)) if image else ""
    page_text = clean_text(soup.get_text(" "))
    profile = company_profile(soup, structured, page_url)
    manufacturer = text_after_label(soup, ("Manufacturer", "Manufactured By")) or profile["company_name"]
    country = text_after_label(soup, ("Country of Origin", "Country Of Origin", "Made In"))
    usage = text_after_label(soup, ("Usage", "Application", "Applications"))
    if usage.lower() in {"industrial", "standard", "general"}:
        usage = ""
    features = unique([clean_text(value) for value in re.findall(r"Features\s+([^:]{1,100})", page_text, re.IGNORECASE)])
    if not features:
        feature_value = text_after_label(soup, ("Features",))
        features = [feature_value] if feature_value else []
    tests = unique([value for value in [usage] if value and any(word in value.lower() for word in ("test", "chemistry", "diagnostic", "analysis"))])
    return {
        "url": page_url,
        "name": title,
        "description": description,
        "image": image,
        "brand": clean_text(str(seller_name)),
        "features": features,
        "applications": [usage] if usage else [],
        "related_products": related_products(soup, page_url),
        "manufacturer_company": clean_text(manufacturer),
        "importer_company": "",
        "country_of_manufacture": clean_text(country),
        "supplier_company": profile["company_name"],
        **profile,
        "brochure_links": [urljoin(page_url, link["href"]) for link in soup.find_all("a", href=True) if ".pdf" in link["href"].lower() or "brochure" in clean_text(link.get_text(" ")).lower()],
        "machine_usage": usage or description,
        "tests_performed": tests,
        "seller_phone": "",
    }


def reveal_seller_phone(url: str, buyer_name: str, buyer_mobile: str, quantity: str = "1") -> str:
    """Reveal a seller number through TradeIndia's own form, if the site allows it."""
    try:
        from playwright.sync_api import sync_playwright
    except ImportError as error:
        raise RuntimeError("Install Playwright first: py -m pip install playwright; py -m playwright install chromium") from error

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=False)
        page = browser.new_page()
        try:
            page.goto(url, wait_until="domcontentloaded", timeout=60_000)
            page.get_by_role("button", name="VIEW NUMBER").click()
            page.locator('input[name="product_name"]').fill(page.title().split(" at ", 1)[0])
            visible_text_inputs = page.locator('input[type="text"]:visible')
            visible_text_inputs.nth(2).fill(buyer_name)
            page.locator('input[name="mobile"]:visible').fill(buyer_mobile)
            page.locator('input[type="number"]:visible').last.fill(quantity)
            page.get_by_role("button", name="CONTINUE").click()
            page.wait_for_timeout(2_000)
            phone_links = page.locator('a[href^="tel:"]:visible')
            for index in range(phone_links.count()):
                href = phone_links.nth(index).get_attribute("href") or ""
                number = re.sub(r"[^+\d]", "", href.removeprefix("tel:"))
                if re.search(r"(?:\+91|91)?[6-9]\d{9}$", number):
                    return number
            phone_match = re.search(r"(?:\+91[\s-]?|91[\s-]?|0)?[6-9]\d{9}", page.locator("body").inner_text())
            if phone_match:
                return clean_text(phone_match.group(0))
            input("No number is visible yet. Complete any TradeIndia verification in the browser, then press Enter here...")
            phone_match = re.search(r"(?:\+91[\s-]?|91[\s-]?|0)?[6-9]\d{9}", page.locator("body").inner_text())
            return clean_text(phone_match.group(0)) if phone_match else ""
        finally:
            browser.close()


def discover_product_urls(search_url: str = SEARCH_URL, timeout: int = 30) -> list[str]:
    pending = [canonical_url(search_url)]
    visited: set[str] = set()
    products: list[str] = []
    while pending:
        page_url = pending.pop(0)
        if page_url in visited:
            continue
        response = requests.get(page_url, headers=HEADERS, timeout=timeout)
        response.raise_for_status()
        visited.add(page_url)
        soup = BeautifulSoup(response.text, "html.parser")
        products.extend(product_links(soup, response.url))
        for next_page in search_page_links(soup, response.url):
            if next_page not in visited and next_page not in pending:
                pending.append(next_page)
    return list(dict.fromkeys(products))


def save_products_to_excel(products: list[dict[str, Any]], output_path: str) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append(COLUMNS)
    for product in products:
        row = dict(product)
        for key in ("features", "applications", "brochure_links", "tests_performed"):
            row[key] = "\n".join(row.get(key, []))
        row["related_products"] = "\n".join(f"{item['name']}: {item['url']}" for item in row.get("related_products", []))
        sheet.append([row.get(column, "") for column in COLUMNS])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.iter_cols(1, len(COLUMNS)):
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = 35
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("url", nargs="?", default=DEFAULT_URL)
    parser.add_argument("-o", "--output", default="tradeindia_products.json")
    parser.add_argument("--excel", default="tradeindia_products.xlsx")
    parser.add_argument("--all", action="store_true", help="Scrape all products from the clinical-analyzer search page")
    parser.add_argument("--delay", type=float, default=0.2)
    parser.add_argument("--reveal-number", action="store_true", help="Use TradeIndia's popup to reveal seller numbers")
    parser.add_argument("--buyer-name", help="Name submitted to TradeIndia's reveal form")
    parser.add_argument("--buyer-mobile", help="Mobile submitted to TradeIndia's reveal form")
    parser.add_argument("--quantity", default="1", help="Quantity submitted to TradeIndia's reveal form")
    args = parser.parse_args()
    if args.reveal_number and (not args.buyer_name or not args.buyer_mobile):
        raise SystemExit("--reveal-number requires --buyer-name and --buyer-mobile; these are submitted to TradeIndia")
    urls = discover_product_urls() if args.all else [canonical_url(args.url)]
    if args.all and canonical_url(args.url) not in urls:
        urls.insert(0, canonical_url(args.url))
    products, errors = [], []
    for index, url in enumerate(dict.fromkeys(urls), 1):
        try:
            print(f"[{index}/{len(urls)}] Scraping {url}")
            product = extract_product(url)
            if args.reveal_number:
                try:
                    product["seller_phone"] = reveal_seller_phone(url, args.buyer_name, args.buyer_mobile, args.quantity)
                except (RuntimeError, ValueError) as error:
                    errors.append({"url": url, "error": str(error)})
            products.append(product)
        except requests.RequestException as error:
            errors.append({"url": url, "error": str(error)})
        if index < len(urls):
            time.sleep(max(args.delay, 0))
    with open(args.output, "w", encoding="utf-8") as file:
        json.dump({"products": products, "errors": errors}, file, indent=2, ensure_ascii=False)
        file.write("\n")
    save_products_to_excel(products, args.excel)
    print(f"Scraped {len(products)} products; {len(errors)} errors")
    print(f"JSON: {args.output}")
    print(f"Excel: {args.excel}")


if __name__ == "__main__":
    main()
